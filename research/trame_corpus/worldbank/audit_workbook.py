"""Read-only audit of the acquired official workbook; no prediction or model fitting."""
import collections
import hashlib
import json
import math
import pathlib
import re
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent
manifest = json.loads((ROOT / 'acquisition_manifest.json').read_text(encoding='utf-8'))
source = ROOT / pathlib.Path(manifest['path']).name
assert hashlib.sha256(source.read_bytes()).hexdigest() == manifest['sha256']
book = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = book['Monthly Prices']
rows = list(sheet.values)
headers, units = rows[4], rows[5]
columns = [i for i, x in enumerate(headers) if i and isinstance(x, str) and x.strip()]
observations = [(rnum, r) for rnum, r in enumerate(rows, 1)
                if isinstance(r[0], str) and re.fullmatch(r'\d{4}M\d{2}', r[0])]
periods = [r[0] for _, r in observations]
duplicates = [p for p, n in collections.Counter(periods).items() if n > 1]
ordinals = [int(p[:4])*12 + int(p[5:])-1 for p in periods]
assert not duplicates
assert ordinals == list(range(ordinals[0], ordinals[-1]+1))
selected_names = ['Gold', 'Silver', 'Copper', 'Crude oil, Brent', 'Maize', 'Wheat, US HRW']
selected = {name: next(i for i in columns if headers[i].strip() == name) for name in selected_names}
stats = {}
missing_tokens = collections.Counter()
for i in columns:
    numeric, missing, invalid = [], [], []
    for rnum, row in observations:
        value = row[i]
        if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
            numeric.append((row[0], value, rnum))
        elif value is None or isinstance(value, str) and (not value.strip() or value.strip() in ('...', '..', '\u2026')):
            missing.append(row[0]); missing_tokens[str(value)] += 1
        else:
            invalid.append({'period': row[0], 'row': rnum, 'value': repr(value)})
    pre2021 = [(p, v, r) for p, v, r in numeric if p <= '2020M12']
    stats[headers[i].strip()] = {
        'column': openpyxl.utils.get_column_letter(i+1), 'unit': units[i],
        'numeric_count': len(numeric), 'missing_count': len(missing),
        'nonpositive_count': sum(v <= 0 for _, v, _ in numeric), 'invalid': invalid,
        'first_numeric_period': numeric[0][0] if numeric else None,
        'last_numeric_period': numeric[-1][0] if numeric else None,
        'count_1960_2020': len(pre2021),
        'missing_1960_2020': sum(p <= '2020M12' for p in missing),
        'source_range_1960_2020': (f'{openpyxl.utils.get_column_letter(i+1)}{pre2021[0][2]}:'
                                   f'{openpyxl.utils.get_column_letter(i+1)}{pre2021[-1][2]}') if pre2021 else None,
    }
gold_description = [(rnum, row[1]) for rnum, row in enumerate(book['Description'].values, 1)
                    if len(row) > 1 and isinstance(row[1], str) and row[1].startswith('Gold,')]
mismatch = list(book['Mismatch Details'].values)
mismatch_rows = [r for r in mismatch[4:] if r[0] is not None]
audit = {
    'kind': 'coverage and semantics audit only; no predictive analysis',
    'source_sha256': manifest['sha256'], 'python': sys.version, 'openpyxl': openpyxl.__version__,
    'workbook_sheets': book.sheetnames, 'updated_annotation': rows[3][0],
    'sheet': sheet.title, 'first_period': periods[0], 'last_period': periods[-1],
    'monthly_rows': len(observations), 'named_series': len(columns),
    'numeric_cells': sum(s['numeric_count'] for s in stats.values()),
    'missing_cells': sum(s['missing_count'] for s in stats.values()),
    'missing_tokens': dict(missing_tokens), 'duplicates': duplicates,
    'calendar_months_contiguous': True, 'unexpected_values': sum(len(s['invalid']) for s in stats.values()),
    'selected_1960_2020_numeric_cells': sum(stats[n]['count_1960_2020'] for n in selected_names),
    'selected_series': {n: stats[n] for n in selected_names},
    'all_series': stats, 'gold_description': gold_description,
    'mismatch_sheet_rows': len(mismatch_rows),
    'mismatch_sheet_interpretation': 'Provider internal comparison, not an authoritative correction instruction; not applied.',
    'limitations': ['All acquired workbook data are exploratory, not an untouched final test.',
                   'A 2026 snapshot does not establish historical availability or vintages.',
                   'Monthly averages and mixed benchmark definitions are not OHLCV bars.',
                   'Composite and related series must not count as independent market histories.',
                   'Quoted first/last periods do not certify stable economic definitions.'],
}
assert audit['numeric_cells'] + audit['missing_cells'] + audit['unexpected_values'] == len(observations)*len(columns)
audit['audit_status'] = 'source exceptions retained explicitly; no numerical replacement'
audit['parser_note'] = ('Initial audit stopped on U+2026 ellipses and four stored Excel #VALUE! error cells. '
                        'Ellipses were then recognized as missing; formula errors remain invalid.')
assert all(audit['selected_series'][n]['missing_1960_2020'] == 0 for n in selected_names)
assert all(not audit['selected_series'][n]['invalid'] for n in selected_names)
(ROOT / 'audit.json').write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding='utf-8')
print(json.dumps({k: v for k, v in audit.items() if k not in ('all_series', 'python')}, indent=2, ensure_ascii=False))
